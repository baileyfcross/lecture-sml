"""PPTX extraction preserving slide boundaries and text-box order."""

from pathlib import Path

from pptx import Presentation

from lecture_slm.ingestion.models import (
    DocumentType,
    ExtractionStatus,
    NormalizedDocument,
    NormalizedSection,
    ProvenanceRecord,
)


class PptxExtractor:
    name = "python-pptx"
    version = "1"
    document_type = DocumentType.PPTX

    def extract(
        self,
        path: Path,
        *,
        source_id: str,
        source_hash: str,
        relative_path: str,
        source_version: int = 1,
    ) -> NormalizedDocument:
        presentation = Presentation(str(path))
        sections: list[NormalizedSection] = []
        warnings: list[str] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            title = slide.shapes.title.text.strip() if slide.shapes.title else None
            body: list[str] = []
            for shape in slide.shapes:
                if shape is slide.shapes.title:
                    continue
                if getattr(shape, "has_table", False):
                    text = "\n".join(
                        " | ".join(cell.text.strip() for cell in row.cells)
                        for row in shape.table.rows
                    ).strip()
                elif getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()
                else:
                    continue
                if text:
                    body.append(text)
            text = "\n".join(body)
            if title or text:
                sections.append(
                    NormalizedSection(
                        section_id=f"slide-{slide_number}",
                        section_type="slide",
                        title=title,
                        order=slide_number - 1,
                        text=text,
                        slide_number=slide_number,
                        source_location=f"slide:{slide_number}",
                        metadata={"title": title, "text_box_count": len(body)},
                    )
                )
            else:
                warnings.append(f"Slide {slide_number} has no extractable text")
        if not sections:
            warnings.append("Presentation has no extractable text")
        return NormalizedDocument(
            source_id=source_id,
            title=path.stem,
            document_type=self.document_type,
            sections=sections,
            extraction_warnings=warnings,
            extraction_quality=(
                ExtractionStatus.SUCCESS_WITH_WARNINGS if warnings else ExtractionStatus.SUCCESS
            ),
            provenance=ProvenanceRecord(
                source_file=relative_path,
                source_hash=source_hash,
                source_version=source_version,
                extraction_method=self.name,
            ),
        )
