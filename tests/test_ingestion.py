"""Offline tests for deterministic ingestion and human review safety."""

import shutil
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from lecture_slm.ingestion.discovery import discover_sources
from lecture_slm.ingestion.export import ExportError, export_approved
from lecture_slm.ingestion.extractors.pdf import PdfExtractor
from lecture_slm.ingestion.importer import import_teaching_materials
from lecture_slm.ingestion.manifest import sha256_file
from lecture_slm.ingestion.models import ReviewStatus
from lecture_slm.ingestion.review import CandidateStore


def write_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Lecture 1 - Logic"
    slide.placeholders[1].text = "First bullet\nSecond bullet"
    table = slide.shapes.add_table(2, 2, 0, 0, 100, 100).table
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Meaning"
    table.cell(1, 0).text = "Logic"
    table.cell(1, 1).text = "Reasoning"
    presentation.save(path)


def write_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Lab 1 - Practice", level=1)
    document.add_paragraph("Try the first exercise.", style="List Bullet")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Input"
    table.rows[0].cells[1].text = "Output"
    document.save(path)


def write_pdf(path: Path) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (Lecture PDF text) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as file:
        writer.write(file)


@pytest.fixture
def source_dir(tmp_path: Path) -> Path:
    root = tmp_path / "materials"
    root.mkdir()
    write_pptx(root / "slides.pptx")
    write_docx(root / "lab.docx")
    write_pdf(root / "lecture.pdf")
    (root / "activity.md").write_text("# Activity\n\nTry this.\n", encoding="utf-8")
    (root / "lecture.txt").write_text("Lecture paragraph.\n\nSecond paragraph.", encoding="utf-8")
    (root / "notes.txt").write_text("General notes without an artifact signal.", encoding="utf-8")
    return root


def test_extractors_incremental_import_and_duplicate_handling(
    source_dir: Path, tmp_path: Path
) -> None:
    data_dir = tmp_path / "data" / "ingestion"
    first = import_teaching_materials(source_dir, data_dir=data_dir, course="cis101", level="intro")
    assert first.supported == 6
    assert first.extracted == 6
    assert first.candidates_created == 5
    assert first.pending_review == 5
    assert any(
        "Meaning" in candidate.expected_output
        for candidate in CandidateStore(data_dir / "candidates.json").candidates.values()
    )

    duplicate = source_dir / "renamed-lecture.txt"
    shutil.copy2(source_dir / "lecture.txt", duplicate)
    second = import_teaching_materials(source_dir, data_dir=data_dir)
    assert second.unchanged == 6
    assert second.duplicates == 1
    assert second.candidates_created == 0
    assert len(CandidateStore(data_dir / "candidates.json").candidates) == 5

    lecture = source_dir / "lecture.txt"
    lecture.write_text("Updated lecture paragraph.", encoding="utf-8")
    third = import_teaching_materials(lecture, data_dir=data_dir, course="cis101", level="intro")
    assert third.extracted == 1
    manifest = (data_dir / "manifests" / "sources.json").read_text(encoding="utf-8")
    assert "previous_source_id" in manifest


def test_blank_pdf_is_success_with_warning_and_not_silent(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as file:
        writer.write(file)
    document = PdfExtractor().extract(
        path,
        source_id="src-" + "a" * 64,
        source_hash=sha256_file(path),
        relative_path="scan.pdf",
    )
    assert document.extraction_warnings
    assert "OCR" in " ".join(document.extraction_warnings)


def test_discovery_refuses_direct_generated_or_evaluation_file(tmp_path: Path) -> None:
    path = tmp_path / "evals" / "prompt.md"
    path.parent.mkdir()
    path.write_text("benchmark", encoding="utf-8")
    with pytest.raises(ValueError, match="generated or evaluation"):
        discover_sources(path)


def test_review_lifecycle_and_approved_only_export(source_dir: Path, tmp_path: Path) -> None:
    data_dir = tmp_path / "data" / "ingestion"
    import_teaching_materials(source_dir, data_dir=data_dir, course="cis101", level="intro")
    store = CandidateStore(data_dir / "candidates.json")
    ids = sorted(store.candidates)
    store.transition(ids[0], ReviewStatus.APPROVED, reviewer="tester", notes="looks good")
    store.transition(ids[1], ReviewStatus.REJECTED, reviewer="tester", notes="not suitable")
    store.transition(ids[2], ReviewStatus.NEEDS_EDIT, reviewer="tester", notes="fix prompt")
    store.transition(ids[3], ReviewStatus.DEFERRED, reviewer="tester")

    output, manifest, count = export_approved(
        data_dir / "candidates.json", tmp_path / "processed", version="0.1.0"
    )
    assert count == 1
    assert output.exists()
    assert manifest.exists()
    assert len(output.read_text(encoding="utf-8").splitlines()) == 1
    exported = CandidateStore(data_dir / "candidates.json").get(ids[0])
    assert exported.review_status is ReviewStatus.APPROVED
    assert exported.quality_tier.value == "A"


def test_export_rejects_approved_candidate_without_required_metadata(
    source_dir: Path, tmp_path: Path
) -> None:
    data_dir = tmp_path / "data" / "ingestion"
    import_teaching_materials(source_dir, data_dir=data_dir)
    store = CandidateStore(data_dir / "candidates.json")
    candidate = next(iter(store.candidates.values()))
    candidate.course = None
    candidate.transition(ReviewStatus.APPROVED, reviewer="tester")
    store.save()
    with pytest.raises(ExportError, match="course"):
        export_approved(data_dir / "candidates.json", tmp_path / "processed", version="0.1.0")
